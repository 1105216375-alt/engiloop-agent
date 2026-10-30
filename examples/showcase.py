"""Three reproducible stories for the EngiLoop homepage.

python -m examples.showcase [report|residual|rules|all] [--json]

Synthetic inputs and scripted candidate text, real extracted runtime/checks.
No LLM, API keys, private archives, or network. This is not a live model demo.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from contextlib import redirect_stdout
from io import BytesIO, StringIO
import json
from pathlib import Path


def report_demo() -> dict:
    """A DOCX draft is checked against separate, explicitly supplied evidence."""
    from docx import Document
    from engineering_agent.report_input import parse_report
    from examples.reverse_verification import synthetic_evidence, synthetic_results

    rows = synthetic_results()
    evidence = synthetic_evidence()
    doc = Document()
    doc.add_heading("合成报告草稿：仅供核验演示", level=1)
    doc.add_paragraph("主梁的风险评分为92。")
    doc.add_paragraph("依据 JTG 0000-2099 第9.9条，应复核记录。")
    table = doc.add_table(rows=2, cols=5)
    for cell, text in zip(table.rows[0].cells, ("序号", "位置", "种类", "具体描述", "照片编号")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, (
        "1", "主梁", "裂缝", "第1跨1号梁腹板斜向裂缝1处，l=0.8m，w=0.3mm", "合成图-01",
    )):
        cell.text = text
    stream = BytesIO()
    doc.save(stream)
    payload = stream.getvalue()
    parsed = parse_report(payload, document_id="synthetic-conflicting-report.docx")
    # Paragraph extraction is deliberately explicit; parse_report handles the table.
    readback = Document(BytesIO(payload))
    candidate = "\n".join(p.text for p in readback.paragraphs[1:])
    blocked = evidence.deliver(candidate, rows)
    wrong_object = evidence.deliver("主梁的风险评分为30。", rows)
    # A template uses the supplied evidence, not the rejected report value.
    corrected_candidate = f"主梁的风险评分为{rows[0]['风险评分']}。"
    checked = evidence.deliver(corrected_candidate, rows)
    return {
        "id": "report", "title": "A convincing report can still be wrong.",
        "synthetic": True, "model_calls": 0,
        "draft": candidate, "evidence": rows,
        "source": {"document_id": parsed["document_id"], "sha256": parsed["sha256"],
                   "records": parsed["records"]},
        "blocked_delivery": blocked, "wrong_object_delivery": wrong_object,
        "corrected_candidate": corrected_candidate, "checked_delivery": checked,
        "boundary": "Separate synthetic evidence is supplied; the report does not validate itself. "
                    "Risk score is not condition score. Scripted correction, not LLM self-repair. "
                    "Covered numeric/citation checks, not proof of arbitrary prose.",
    }


def residual_demo() -> dict:
    """Add evidence to one fixture session and re-run the actual reasoning mode."""
    from engineering_agent.bridge_fixture import (
        build_bridge_fixture, analyze_fixture, unknown_keys, accepted_action_types,
    )
    from engineering_agent.nlp import MeasurementIntakeSession
    from examples.checkpoints_and_rollback import run_runtime_demo

    fixture = build_bridge_fixture(dr=None, bci=None)
    session = fixture.session_store.get_or_create(fixture.session_id)
    stages = []
    for label, added in (("missing_both", None), ("dr_added", "Dr"), ("bci_added", "BCI")):
        if added == "Dr":
            session.metadata["bridge_score"]["全桥评分"] = 92
        elif added == "BCI":
            session.metadata["bridge_score_52xx"]["BCI"] = 94
        result = analyze_fixture(fixture, "整体风险如何？")
        remaining = sorted(unknown_keys(result))
        stages.append({
            "stage": label, "added_synthetic_input": added,
            "remaining_requirements": remaining, "reply": result["reply"],
            "accepted_actions": accepted_action_types(result),
            "runtime_status": result["final_state_status"],
            "cross_indicator_screening_present": "技术状况风险筛查:" in result["reply"],
            "engineering_signoff": False,
        })
    intake = MeasurementIntakeSession()
    turns = ["1-1号T梁梁底有一处锈胀露筋。",
             "距0号台6米、距右侧0米，一巴掌长、两指宽。",
             "尺子量过：长1.5米，宽十厘米。"]
    intake_steps = [{"input": text, "result": deepcopy(intake.ingest(text))} for text in turns]
    return {
        "id": "residual", "title": "Missing evidence is unfinished work.",
        "synthetic": True, "model_calls": 0, "stages": stages,
        "separate_intake_example": intake_steps,
        "separate_fault_injection": run_runtime_demo(),
        "boundary": "Scores are supplied test inputs, not computed from the intake. "
                    "Each stage re-evaluates the updated session, not exact mid-step resume. "
                    "Runtime 'solved' may mean a bounded answer; it does not clear unknowns. "
                    "Intake and rollback are separately labeled mechanism examples.",
    }


def rules_demo() -> dict:
    """A versioned data mapping changes deductions; the scoring engine is unchanged."""
    from engineering_agent.rulepack import evaluate_rulepack, snapshot_reuse_status

    directory = Path(__file__).parent / "fixtures"
    packs = [json.loads((directory / f"rules_v{n}.json").read_text(encoding="utf-8"))
             for n in (1, 2)]
    context = {
        "observations": [
            {"构件": "1号主梁", "构件类别": "主梁", "病害": "合成裂缝", "标度": 3,
             "source": "synthetic://observations/main-1"},
            {"构件": "1号支座", "构件类别": "支座", "病害": "合成支座脱空", "标度": 2,
             "source": "synthetic://observations/bearing-1"},
        ],
        "selected_parts": ["upper/main", "upper/bear"],
        "structure": {"总主梁数": 1, "总支座数": 1},
        "part_applicability": {"upper": {
            "main": {"applicable": True, "source": "synthetic://scope/main"},
            "bear": {"applicable": True, "source": "synthetic://scope/bearing"},
        }},
    }
    before = evaluate_rulepack(packs[0], **context)
    unchanged = snapshot_reuse_status(before, packs[0], context=context)
    stale = snapshot_reuse_status(before, packs[1], context=context)
    after = evaluate_rulepack(packs[1], **context)
    current = snapshot_reuse_status(after, packs[1], context=context)
    return {
        "id": "rules", "title": "Update the rules. Not the model.",
        "synthetic": True, "model_calls": 0,
        "before": before, "after": after, "unchanged_context": unchanged,
        "old_result_under_new_rules": stale, "recomputed_context": current,
        "boundary": "Invented rules, not JTG provisions. Only the explicit mapping changes. "
                    "Rule/source hashes track identity, not truth or authorization. "
                    "Reuse checks are explicit calls, not a production invalidation service.",
    }


def run_showcase(which: str = "all") -> dict:
    runners = {"report": report_demo, "residual": residual_demo, "rules": rules_demo}
    if which not in {*runners, "all"}:
        raise ValueError("Unknown demonstration")
    # Legacy imports emit an ontology banner. Keep machine output valid JSON.
    with redirect_stdout(StringIO()):
        results = [run() for name, run in runners.items() if which in ("all", name)]
    return {"schema": 1, "synthetic": True, "model_calls": 0,
            "notice": "Executable mechanism demos, not a live model or certified assessment.",
            "demos": results}


def summary(result: dict) -> str:
    lines = ["EngiLoop | Verifiable Domain Agents", "Synthetic inputs. Real execution. No model calls.", ""]
    for demo in result["demos"]:
        lines.append(demo["title"])
        if demo["id"] == "report":
            lines.extend([
                f"  Report claims risk=92; separate main-girder evidence={demo['evidence'][0]['风险评分']}.",
                f"  Wrong number + missing clause -> {demo['blocked_delivery']['status'].upper()}",
                f"  Bearing value used for girder -> {demo['wrong_object_delivery']['status'].upper()}",
                f"  Evidence-bound replacement -> {demo['checked_delivery']['status'].upper()} (mechanical checks)",
            ])
        elif demo["id"] == "residual":
            for stage in demo["stages"]:
                lines.append(f"  {stage['stage']}: remaining={stage['remaining_requirements']}; "
                             f"screening={stage['cross_indicator_screening_present']}")
            gaps = [len(s["result"]["missing_fields"]) for s in demo["separate_intake_example"]]
            lines.append(f"  Separate intake: missing fields {gaps}; source review still required.")
            rollback = demo["separate_fault_injection"]["rollback_after_injected_invalid_update"]
            lines.append(f"  Separate fault injection: rollbacks={rollback['rollbacks']}; "
                         f"old state preserved={rollback['old_state_preserved']}")
        else:
            for key in ("before", "after"):
                run = demo[key]
                computed = run["snapshot"]["calculated"]
                lines.append(f"  Rule pack {run['binding']['version']}: "
                             f"main={computed['component_scores']['upper/main/1']}; "
                             f"selected-scope={computed['selected_scope_score']}")
            lines.append(f"  Old snapshot reusable under v2: {demo['old_result_under_new_rules']['reusable']}")
            lines.append(f"  Recomputed snapshot context matches: {demo['recomputed_context']['reusable']}")
        lines.append("")
    lines.append("No automatic engineering sign-off. Use --json for evidence, issues, state and source details.")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("demo", nargs="?", choices=("all", "report", "residual", "rules"), default="all")
    parser.add_argument("--json", action="store_true", help="Emit full machine-readable evidence")
    args = parser.parse_args()
    result = run_showcase(args.demo)
    print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else summary(result))


if __name__ == "__main__":
    main()

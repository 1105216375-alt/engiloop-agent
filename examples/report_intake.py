"""Build a synthetic DOCX entirely in memory, then run the real table parser.

Run from the repository root: python -m examples.report_intake
"""
import json
from io import BytesIO

from docx import Document

from engineering_agent.report_input import parse_report


def synthetic_report_bytes() -> bytes:
    doc = Document()
    doc.add_heading("合成桥梁报告：仅供开源机制测试", level=1)
    table = doc.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells, ("序号", "位置", "种类", "具体描述", "照片编号")):
        cell.text = text
    rows = [
        ("1", "主梁", "裂缝", "第1跨1号梁梁底裂缝1处，l=0.8m，w=0.3mm", "合成图-01"),
        ("2", "主梁", "锈胀露筋", "1-1号T梁梁底，距0号台6m，锈胀露筋1处，S=1.5m×0.1m", "合成图-02"),
        ("3", "支座", "支座脱空", "第1跨1号支座脱空已修复，供检查跳过逻辑", "合成图-03"),
    ]
    for values in rows:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def main():
    result = parse_report(synthetic_report_bytes(), document_id="synthetic-bridge.docx")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

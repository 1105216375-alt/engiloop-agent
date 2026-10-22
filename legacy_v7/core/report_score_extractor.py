"""Extract and reconcile condition scores explicitly published in bridge reports.

The disease scorer and the report's signed assessment are deliberately kept as
two independent tracks.  A published score is the delivery record; a score
recomputed from extracted defects is a verification result.  They must never be
silently substituted for one another.
"""

from __future__ import annotations

import copy
import re
from collections import Counter
from typing import Any, Dict, Iterable, List, Optional, Tuple


SCORE_CONTRACT_VERSION = "bridge_score_provenance.v1"
SCORE_CALIBRATION_VERSION = "bridge_score_calibration.v2"
REVIEW_GAP_THRESHOLD = 2.0
CONSISTENCY_TOLERANCE = 0.1

SYSTEM_NAMES = ("上部结构", "下部结构", "桥面系")
SYSTEM_FORMULA_WEIGHTS = {"桥面系": 0.2, "上部结构": 0.4, "下部结构": 0.4}
SYSTEM_INTERNAL_KEYS = {"上部结构": "upper", "下部结构": "lower", "桥面系": "deck"}

PART_ALIASES = (
    (("上部承重构件", "主要承重构件", "主承重构件"), "main"),
    (("上部一般构件", "一般构件"), "gen"),
    (("伸缩缝装置", "伸缩装置", "伸缩缝"), "joints"),
    (("栏杆、护栏", "栏杆护栏", "防撞护栏", "护栏"), "railing"),
    (("照明、标志", "照明标志", "标志标线", "照明"), "signs"),
    (("桥面铺装", "桥面铺装层"), "pavement"),
    (("排水系统", "桥面排水"), "drain"),
    (("人行道", "人行道、路缘石"), "sidewalk"),
    (("墩台基础", "桥墩基础", "桥台基础", "基础"), "found"),
    (("翼墙、耳墙", "翼墙耳墙", "翼墙", "耳墙"), "wing"),
    (("锥坡、护坡", "锥坡护坡", "锥坡", "护坡"), "cone"),
    (("调治构造物", "调治构造"), "reg"),
    (("桥路连接处", "桥头搭板"), "approach"),
    (("桥墩", "墩柱", "盖梁"), "pier"),
    (("桥台",), "abut"),
    (("支座",), "bear"),
    (("河床",), "river"),
)


def _clean_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def _row_values(row: Any) -> List[str]:
    values: List[str] = []
    for cell in row.cells:
        value = _clean_text(cell.text)
        if value and (not values or values[-1] != value):
            values.append(value)
    return values


def _as_float(value: Any) -> Optional[float]:
    text = str(value or "").strip().replace("％", "%")
    match = re.fullmatch(r"-?\d+(?:\.\d+)?%?", text)
    if not match:
        return None
    number = float(text.rstrip("%"))
    if text.endswith("%"):
        number /= 100.0
    return number


def _optional_float(value: Any) -> Optional[float]:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def build_score_calibration_audit(
    reported_dr: Any,
    computed_dr: Any,
    *,
    reported_support: int = 0,
    reported_source_locator: str = "",
    computed_reliability: Optional[Dict[str, Any]] = None,
    total_records: int = 0,
    diagnosed_records: int = 0,
    coverage: Optional[Dict[str, Any]] = None,
    calibration_evidence: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Explain a report/recomputed score gap without mutating either score.

    Gap bands are audit priorities, not bridge condition grades.  Even a proven
    input defect does not authorize automatic score replacement; manual
    calibration additionally requires equivalent inputs, equivalent formulas,
    and reviewer approval.
    """
    reported = _optional_float(reported_dr)
    computed = _optional_float(computed_dr)
    reliability = dict(computed_reliability or {})
    coverage = dict(coverage or {})
    calibration_evidence = dict(calibration_evidence or {})
    causes: List[Dict[str, Any]] = []
    actions: List[str] = []

    gap = round(computed - reported, 2) if reported is not None and computed is not None else None
    absolute_gap = round(abs(gap), 2) if gap is not None else None
    if gap is None:
        status, severity = "not_comparable", "info"
        direction = "unavailable"
    elif absolute_gap <= 1.0:
        status, severity = "aligned", "ok"
        direction = "aligned"
    elif absolute_gap <= 3.0:
        status, severity = "minor_gap", "info"
        direction = "system_higher" if gap > 0 else "system_lower"
    elif absolute_gap <= 8.0:
        status, severity = "review_required", "warning"
        direction = "system_higher" if gap > 0 else "system_lower"
    else:
        status, severity = "severe_conflict", "critical"
        direction = "system_higher" if gap > 0 else "system_lower"

    def add_cause(code: str, label: str, basis: str, *, proven: bool, role: str) -> None:
        causes.append({
            "code": code,
            "label": label,
            "basis": basis,
            "proven": bool(proven),
            "certainty": "已证实" if proven else "待核可能",
            "role": role,
        })

    if gap is None:
        missing = []
        if reported is None:
            missing.append("报告Dr")
        if computed is None:
            missing.append("系统重算Dr")
        add_cause(
            "MISSING_COMPARABLE_SCORE",
            "缺少可比较评分",
            "、".join(missing) + "未形成",
            proven=True,
            role="comparison_blocker",
        )
        actions.append("补齐缺失评分及来源定位后再进行差值审计。")
    else:
        missing_fields = list(reliability.get("missing_structure_fields") or [])
        if reliability.get("publishable") is False or reliability.get("provisional") is True:
            basis = "系统重算被标记为暂定值"
            if missing_fields:
                basis += "；缺少" + "、".join(str(row) for row in missing_fields[:8])
            add_cause(
                "COMPUTED_INPUT_INCOMPLETE",
                "系统重算输入不完整",
                basis,
                proven=True,
                role="evidence_quality",
            )
            causes[-1]["proof_scope"] = "input_quality_only_not_gap_causation"
            causes[-1]["basis"] += "；这只证明本轮重算有待核条件，不能据此确证分差成因"
            actions.append("核对病害定量条件、评分规则覆盖及关键构件参数，完成待核输入后按同一版本规则复算。")

        report_traceable = bool(str(reported_source_locator or "").strip()) or int(reported_support or 0) > 0
        if reported is not None and not report_traceable:
            add_cause(
                "REPORT_SOURCE_TRACEABILITY_GAP",
                "报告评分定位证据不足",
                "已提取报告Dr，但未保留评分表定位或支持记录",
                proven=True,
                role="evidence_quality",
            )
            actions.append("回到报告正式评定表，补录页码、表名和评分摘录。")

        truncated = bool(coverage.get("truncated"))
        incomplete_records = bool(total_records and diagnosed_records < total_records)
        if truncated or incomplete_records:
            basis = "证据契约发生截断" if truncated else ""
            if incomplete_records:
                record_basis = f"仅诊断{diagnosed_records}/{total_records}条记录"
                basis = (basis + "；" + record_basis).strip("；")
            add_cause(
                "EXTRACTION_COVERAGE_LIMITED",
                "病害抽取覆盖不完整",
                basis,
                proven=True,
                role="evidence_quality",
            )
            causes[-1]["proof_scope"] = "coverage_quality_only_not_gap_causation"
            causes[-1]["basis"] += "；覆盖限制已知，但尚未证明其造成当前分差"
            actions.append("完成全量病害抽取并核对漏项后再复算。")

        if status == "aligned":
            add_cause(
                "ROUNDING_OR_AGGREGATION",
                "舍入或汇总精度差异",
                f"绝对差值{absolute_gap}分，处于1分审计容差内",
                proven=False,
                role="possible_gap_driver",
            )
            actions.append("保留双轨来源，无需调整任何已发布评分。")
        else:
            add_cause(
                "FORMULA_SCOPE_MISMATCH",
                "评分公式或适用构件范围不同",
                "需逐项核对报告适用权重、系统构件映射和扣分聚合路径",
                proven=False,
                role="possible_gap_driver",
            )
            if absolute_gap is not None and absolute_gap > 8.0:
                add_cause(
                    "COMPONENT_MAPPING_OR_WEIGHT_SCOPE",
                    "构件映射或权重适用范围可能不同",
                    f"绝对差值{absolute_gap}分已超过严重冲突阈值8分",
                    proven=False,
                    role="possible_gap_driver",
                )
            actions.append("逐项对齐报告部件权重、病害构件归属、适用范围和扣分链。")
            actions.append("原因未证明前，报告正式值与系统重算值继续双轨展示。")

    # Known input/coverage defects prove a reliability limitation, not the
    # counterfactual contribution of that limitation to a score difference.
    # This path has no audited attribution/reconciliation executor, therefore
    # none of the quality findings above can satisfy the causal-proof gate.
    proven_gap_driver = any(
        row.get("proven") and row.get("role") == "gap_driver" for row in causes
    )
    gate_conditions = {
        "gap_cause_proven": proven_gap_driver,
        "inputs_equivalent": bool(calibration_evidence.get("inputs_equivalent")),
        "formulas_equivalent": bool(calibration_evidence.get("formulas_equivalent")),
        "reviewer_approved": bool(calibration_evidence.get("reviewer_approved")),
    }
    calibration_allowed = bool(gap is not None and status != "aligned" and all(gate_conditions.values()))
    if gap is not None and status != "aligned":
        actions.append("先逐项证实分差成因，且同源输入、同一公式和人工审批均留痕后，才允许人工校准。")

    diagnostic_confidence = "high" if proven_gap_driver else ("medium" if gap is not None else "high")
    return {
        "audit_version": SCORE_CALIBRATION_VERSION,
        "status": status,
        "severity": severity,
        "reported_dr": round(reported, 2) if reported is not None else None,
        "computed_dr": round(computed, 2) if computed is not None else None,
        "gap": gap,
        "gap_computed_minus_reported": gap,
        "absolute_gap": absolute_gap,
        "direction": direction,
        "band_policy": "<=1 aligned; <=3 minor; <=8 review; >8 severe",
        "band_scope": "差值审计优先级，不是桥梁技术状况等级",
        "primary_track": "report_conclusion" if reported is not None else "system_recomputed_fallback",
        "decision": "approved_manual_calibration" if calibration_allowed else "keep_dual_track",
        "no_overwrite": bool(reported is not None and not calibration_allowed),
        "calibration_allowed": calibration_allowed,
        "automatic_calibration_allowed": False,
        "diagnostic_confidence": diagnostic_confidence,
        "gap_cause_status": "proven" if proven_gap_driver else "not_proven",
        "quality_findings_do_not_prove_gap_cause": True,
        "causes": causes,
        "required_actions": list(dict.fromkeys(actions)),
        "calibration_gate": {
            "conditions": gate_conditions,
            "passed": calibration_allowed,
            "policy": "自动校准永久禁止；人工校准需原因已证实、输入同源、公式一致并经审批。",
        },
    }


def _system_name(value: str) -> str:
    clean = _clean_text(value)
    for name in SYSTEM_NAMES:
        if name in clean:
            return name
    return ""


def _part_key(value: str) -> str:
    clean = _clean_text(value)
    for aliases, key in PART_ALIASES:
        if any(alias in clean for alias in aliases):
            return key
    return ""


def _is_not_applicable(value: str) -> bool:
    clean = _clean_text(value)
    return clean in {"--", "—", "-", "/", "不适用", "无"}


def extract_part_applicability(reported_score: Optional[Dict[str, Any]]) -> Dict[str, Dict[str, bool]]:
    """Return only the report's applicable/not-applicable topology.

    Published scores remain on their independent evidence track.  The
    recomputation may reuse the report's component inventory because a missing
    disease row means "observed healthy", while an explicit ``--`` means the
    component is not installed.  No published score or weight is copied.
    """
    applicability: Dict[str, Dict[str, bool]] = {}
    for system_name, parts in (reported_score or {}).get("parts", {}).items():
        system_key = SYSTEM_INTERNAL_KEYS.get(system_name)
        if not system_key or not isinstance(parts, dict):
            continue
        for part_key, row in parts.items():
            if not isinstance(row, dict) or "applicable" not in row:
                continue
            applicability.setdefault(system_key, {})[str(part_key)] = bool(row["applicable"])
    return applicability


def _grade_candidates(text: str) -> Iterable[str]:
    patterns = (
        r"技术状况等级(?:评定)?\s*(?:为|：|:|=|＝)\s*([一二三四五1-5ⅠⅡⅢⅣⅤIVXivx]+类)",
        r"评定为\s*([一二三四五1-5ⅠⅡⅢⅣⅤIVXivx]+类)",
    )
    for pattern in patterns:
        for match in re.finditer(pattern, text, re.I):
            yield match.group(1)


def _extract_summary_row(values: List[str]) -> Optional[Dict[str, Any]]:
    system_index = next(
        (index for index, value in enumerate(values) if _system_name(value)),
        None,
    )
    if system_index is None:
        return None
    system = _system_name(values[system_index])

    part_index = next(
        (
            index
            for index in range(system_index + 1, len(values))
            if _part_key(values[index])
        ),
        None,
    )
    if part_index is None:
        return None

    trailing = values[part_index + 1 :]
    weight_index = next(
        (
            index
            for index, value in enumerate(trailing)
            if _is_not_applicable(value)
            or (_as_float(value) is not None and 0 <= float(_as_float(value)) <= 1)
        ),
        None,
    )
    if weight_index is None:
        return None

    weight_token = trailing[weight_index]
    applicable = not _is_not_applicable(weight_token)
    weight = _as_float(weight_token) if applicable else None
    score_values = [
        number
        for number in (_as_float(value) for value in trailing[weight_index + 1 :])
        if number is not None and 0 <= number <= 100
    ]

    # Merged Word cells can collapse identical part/system scores into one value.
    if applicable and not score_values:
        return None
    if not applicable and not score_values:
        return None

    return {
        "system": system,
        "part_key": _part_key(values[part_index]),
        "part_label": values[part_index],
        "applicable": applicable,
        "weight": round(float(weight), 6) if weight is not None else None,
        "score": round(float(score_values[0]), 2) if applicable else None,
        "system_score": round(float(score_values[-1]), 2),
    }


def extract_reported_condition_score(docx_path: str) -> Dict[str, Any]:
    """Return the report-published Dr, grade, system scores and part scores."""
    try:
        from docx import Document

        document = Document(docx_path)
    except Exception as exc:
        return {
            "found": False,
            "contract_version": SCORE_CONTRACT_VERSION,
            "error": str(exc)[:200],
        }

    fragments: List[Tuple[str, str]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            fragments.append((f"paragraph:{index + 1}", text))

    summary_rows: List[Dict[str, Any]] = []
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            values = _row_values(row)
            if not values:
                continue
            locator = f"table:{table_index + 1}/row:{row_index + 1}"
            fragments.append((locator, " | ".join(values)))
            summary = _extract_summary_row(values)
            if summary:
                summary["source_locator"] = locator
                summary_rows.append(summary)

    dr_candidates: List[Dict[str, Any]] = []
    dr_pattern = re.compile(
        r"(?:评分\s*)?Dr\s*值?\s*(?:为|：|:|=|＝)\s*(\d{1,3}(?:\.\d+)?)",
        re.I,
    )
    for locator, text in fragments:
        grades = list(_grade_candidates(text))
        for match in dr_pattern.finditer(text):
            value = float(match.group(1))
            if not 0 <= value <= 100:
                continue
            excerpt = text[max(0, match.start() - 80) : match.end() + 180]
            dr_candidates.append(
                {
                    "dr": round(value, 2),
                    "grade": grades[0] if grades else "",
                    "source_locator": locator,
                    "source_excerpt": excerpt[:280],
                }
            )

    selected_dr: Dict[str, Any] = {}
    support_count = 0
    if dr_candidates:
        counts = Counter((item["dr"], item["grade"]) for item in dr_candidates)
        (dr_value, grade_value), support_count = counts.most_common(1)[0]
        selected_dr = next(
            item
            for item in dr_candidates
            if item["dr"] == dr_value and item["grade"] == grade_value
        )

    systems: Dict[str, float] = {}
    parts: Dict[str, Dict[str, Dict[str, Any]]] = {}
    source_locators: List[str] = []
    for row in summary_rows:
        system = row["system"]
        systems[system] = row["system_score"]
        parts.setdefault(system, {})[row["part_key"]] = {
            "label": row["part_label"],
            "applicable": row["applicable"],
            "weight": row["weight"],
            "score": row["score"],
            "source_locator": row["source_locator"],
        }
        source_locators.append(row["source_locator"])

    formula_check: Dict[str, Any] = {"available": False}
    if all(system in systems for system in SYSTEM_FORMULA_WEIGHTS):
        recomputed = sum(
            systems[system] * weight
            for system, weight in SYSTEM_FORMULA_WEIGHTS.items()
        )
        formula_check = {
            "available": True,
            "formula": "Dr=0.2×桥面系+0.4×上部结构+0.4×下部结构",
            "recomputed_from_rounded_system_scores": round(recomputed, 2),
            "published_dr": selected_dr.get("dr"),
            "rounding_gap": (
                round(recomputed - float(selected_dr["dr"]), 2)
                if selected_dr.get("dr") is not None
                else None
            ),
        }

    found = bool(selected_dr or summary_rows)
    return {
        "found": found,
        "contract_version": SCORE_CONTRACT_VERSION,
        "source_type": "report_published" if found else "not_found",
        "dr": selected_dr.get("dr"),
        "grade": selected_dr.get("grade", ""),
        "support_count": support_count,
        "source_locator": selected_dr.get("source_locator", ""),
        "source_excerpt": selected_dr.get("source_excerpt", ""),
        "has_summary_table": bool(summary_rows),
        "summary_row_count": len(summary_rows),
        "summary_source_locators": list(dict.fromkeys(source_locators)),
        "systems": systems,
        "parts": parts,
        "formula_check": formula_check,
    }


def reconcile_condition_scores(
    recomputed: Dict[str, Any], reported: Dict[str, Any]
) -> Dict[str, Any]:
    """Build the primary score and an explicit two-track provenance contract."""
    recomputed = recomputed or {}
    reported = reported or {"found": False}
    primary = copy.deepcopy(recomputed)

    computed_dr = recomputed.get("全桥评分", recomputed.get("Dr"))
    computed_grade = recomputed.get("技术状况等级", recomputed.get("等级", ""))
    reported_dr = reported.get("dr") if reported.get("found") else None
    reported_grade = reported.get("grade", "") if reported.get("found") else ""
    computed_reliability = copy.deepcopy(recomputed.get("评分可靠性") or {})
    computed_publishable = computed_reliability.get("publishable", True) is not False
    structure_review_required = not computed_publishable

    gap = None
    if reported_dr is not None and computed_dr is not None:
        gap = round(float(computed_dr) - float(reported_dr), 2)

    if reported_dr is not None:
        status = (
            "requires_review"
            if structure_review_required or (gap is not None and abs(gap) >= REVIEW_GAP_THRESHOLD)
            else "consistent"
        )
        primary_source = "report_published"
        score_role = "report_primary_system_audit"
        system_score_role = "audit_only"
    else:
        status = "recomputed_requires_structure_review" if structure_review_required else "recomputed_only"
        primary_source = "system_recomputed"
        score_role = "system_fallback_missing_report_score"
        system_score_role = "fallback_estimate"

    review_inputs = []
    missing_structure = list(computed_reliability.get("missing_structure_fields") or [])
    uncertain_grading = int(computed_reliability.get("uncertain_grading_count") or 0)
    applicability_conflicts = list(computed_reliability.get("applicability_conflicts") or [])
    if missing_structure:
        review_inputs.append("关键构件数量:" + "、".join(str(value) for value in missing_structure))
    if uncertain_grading:
        review_inputs.append(f"定量标度:{uncertain_grading}条")
    if applicability_conflicts:
        review_inputs.append(f"适用性冲突:{len(applicability_conflicts)}项")

    if structure_review_required:
        pending = "；".join(review_inputs) or "评分输入待复核"
        reason = f"系统重算存在待补证据({pending})，仅作暂定校核值"
    elif status == "requires_review":
        reason = "报告发布值与系统按抽取病害重算值口径不同"
    elif status == "consistent":
        reason = "报告发布值与系统重算值在复核阈值内"
    else:
        reason = "报告未发布可提取的Dr，采用系统重算值"

    provenance = {
        "contract_version": SCORE_CONTRACT_VERSION,
        "primary_source": primary_source,
        "score_role": score_role,
        "report_score_present": reported_dr is not None,
        "system_score_role": system_score_role,
        "primary_value_policy": (
            "keep_report_value_and_audit_recomputed"
            if reported_dr is not None
            else "use_recomputed_only_when_report_score_missing"
        ),
        "status": status,
        "reported_dr": reported_dr,
        "reported_grade": reported_grade,
        "computed_dr": computed_dr,
        "computed_grade": computed_grade,
        "computed_publishable": computed_publishable,
        "computed_provisional": structure_review_required,
        "computed_reliability": computed_reliability,
        "gap_computed_minus_reported": gap,
        "consistency_tolerance": CONSISTENCY_TOLERANCE,
        "review_gap_threshold": REVIEW_GAP_THRESHOLD,
        "requires_review": status in {"requires_review", "recomputed_requires_structure_review"},
        "reason": reason,
        "reported_source_locator": reported.get("source_locator", ""),
        "reported_summary_table": bool(reported.get("has_summary_table")),
    }

    if reported_dr is not None:
        primary["全桥评分"] = reported_dr
        if reported_grade:
            primary["技术状况等级"] = reported_grade
        if reported.get("systems"):
            primary["系统得分"] = copy.deepcopy(reported["systems"])
        official_part_scores = {
            system: {
                part_key: part["score"]
                for part_key, part in system_parts.items()
                if part.get("score") is not None
            }
            for system, system_parts in (reported.get("parts") or {}).items()
        }
        official_part_scores = {
            system: scores for system, scores in official_part_scores.items() if scores
        }
        if official_part_scores:
            primary["部位得分"] = official_part_scores
        primary["summary"] = (
            f"报告正式评定 Dr={reported_dr}，等级={reported_grade or computed_grade or '未载明'}；"
            f"系统重算 Dr={computed_dr}，差值={gap if gap is not None else '不可计算'}分"
        )

    primary["评分来源"] = "报告正式评定表（主值）" if reported_dr is not None else "系统缺分补算（待复核）"
    primary["评分来源代码"] = primary_source
    primary["评分角色"] = score_role
    primary["评分口径审计"] = provenance
    primary["评分可靠性"] = computed_reliability
    primary["报告适用权重"] = {
        system: {
            part_key: part.get("weight")
            for part_key, part in system_parts.items()
            if part.get("applicable")
        }
        for system, system_parts in (reported.get("parts") or {}).items()
    }
    primary["系统重算摘要"] = {
        "全桥评分": computed_dr,
        "技术状况等级": computed_grade,
        "系统得分": copy.deepcopy(recomputed.get("系统得分") or {}),
        "部位得分": copy.deepcopy(recomputed.get("部位得分") or {}),
    }

    return {
        "primary": primary,
        "reported": reported,
        "recomputed": recomputed,
        "provenance": provenance,
    }

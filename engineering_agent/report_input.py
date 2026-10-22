"""New row-provenance adapter around the production Word-table parser.

The defect parser is unchanged. Each original row is passed through it as an
in-memory table so duplicate descriptions cannot acquire the wrong row source.
This module never loads application configuration or production report paths.
"""
from __future__ import annotations

import hashlib
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document

from legacy_v7.core.report_parser_v2 import (
    _find_v2_columns, _is_v2_table, extract_defects_v2,
)


def parse_report(source: str | Path | bytes, *, document_id: str | None = None) -> dict[str, Any]:
    """Parse a supported DOCX table and bind each emitted row to its source.

    Coordinates are one-based Word table/row/column indices. A content hash is
    an identity/integrity aid, not proof of the report's factual correctness.
    Unsupported tables are counted, never silently represented as understood.
    """
    payload = source if isinstance(source, bytes) else Path(source).read_bytes()
    if len(payload) > 25 * 1024 * 1024:
        raise ValueError("This public adapter accepts DOCX inputs up to 25 MiB.")
    with zipfile.ZipFile(BytesIO(payload)) as package:
        if sum(info.file_size for info in package.infolist()) > 100 * 1024 * 1024:
            raise ValueError("The expanded DOCX package exceeds 100 MiB.")
    content_hash = hashlib.sha256(payload).hexdigest()
    name = document_id or ("in-memory.docx" if isinstance(source, bytes) else Path(source).name)
    doc = Document(BytesIO(payload))
    records: list[dict[str, Any]] = []
    supported = 0
    skipped_rows = 0
    for table_index, table in enumerate(doc.tables, start=1):
        if not _is_v2_table(table):
            continue
        supported += 1
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        col_pos, col_kind, col_desc, col_len, col_wid = _find_v2_columns(headers)
        for row_index, row in enumerate(table.rows[1:], start=2):
            values = [cell.text for cell in row.cells]
            isolated = Document()
            one = isolated.add_table(rows=2, cols=max(len(headers), len(values)))
            for index, value in enumerate(headers):
                one.cell(0, index).text = value
            for index, value in enumerate(values):
                one.cell(1, index).text = value
            buffer = BytesIO()
            isolated.save(buffer)
            buffer.seek(0)
            parsed = extract_defects_v2(buffer)
            if not parsed:
                skipped_rows += 1
                continue
            for item in parsed:
                source_record = {
                    "document_id": name, "sha256": content_hash,
                    "table": table_index, "row": row_index,
                    "description_column": col_desc + 1,
                    "description_text": values[col_desc].strip().replace("\n", " "),
                    "candidate_source_columns": {
                        "component": [col_pos + 1, col_desc + 1],
                        "disease": [col_kind + 1, col_desc + 1],
                        "description": [col_desc + 1]},
                }
                for field, column in (("length", col_len), ("width", col_wid)):
                    if column >= 0 and column < len(values) and values[column].strip():
                        source_record["candidate_source_columns"][field] = [column + 1, col_desc + 1]
                records.append({
                    **item, "source": source_record,
                    "engineering_verified": False, "eligible_for_scoring": False,
                    "scale_is_reported_fact": bool(item.get("_scale_explicit")),
                    "review_status": "待核对报告来源、单位及估计字段",
                })
    return {
        "adapter": "new_public_row_provenance_v1",
        "parser": "production_report_parser_v2",
        "document_id": name, "sha256": content_hash,
        "table_count": len(doc.tables), "supported_table_count": supported,
        "unsupported_table_count": len(doc.tables) - supported,
        "skipped_rows_in_supported_tables": skipped_rows,
        "records": records,
        "boundary": "结构化不等于事实确证；估计标度不是原始报告明示标度。",
    }
